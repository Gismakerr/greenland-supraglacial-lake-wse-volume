from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

import geopandas as gpd
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(r"X:\2024_西南水位曲线")
RESULT1_ROOT = PROJECT_ROOT / "result1"
RESULT2_ROOT = PROJECT_ROOT / "result2"
OUTPUT_DIR = RESULT2_ROOT / "09_1_lake_observation_status_records" / "no_noise_lakes"

S2_STATUS_WORKBOOK = OUTPUT_DIR / "S2统计_无噪声固定水期.xlsx"
SWOT_STATUS_WORKBOOK = OUTPUT_DIR / "SWOT统计_无噪声固定水期.xlsx"
LEGACY_DAILY_WORKBOOK = OUTPUT_DIR / "SWOT按日统计_无噪声固定水期.xlsx"
SUMMARY_WORKBOOK = OUTPUT_DIR / "无噪声固定水期汇总.xlsx"
RESUMMARY_WORKBOOK = OUTPUT_DIR / "无噪声固定水期再汇总.xlsx"
STATUS_WORKBOOK = SWOT_STATUS_WORKBOOK

S2_START = pd.Timestamp("2024-06-05")
S2_END = pd.Timestamp("2024-08-15")

S2_PREVIEW_ROOT = RESULT1_ROOT / "03_Classified_Previews"
S2_NDSI_DIR = RESULT1_ROOT / "01_Preprocessing_Results" / "02_NDSI"
S2_CLOUD_DIR = RESULT1_ROOT / "01_Preprocessing_Results" / "06_Cloud_Mask"
S2_DAILY_ALL_CSV = S2_PREVIEW_ROOT / "Daily_Image_Classification_Summary_All.csv"
S2_POLLUTED_ALL_CSV = S2_PREVIEW_ROOT / "Polluted_Image_Details_All.csv"
S2_EFFECTIVE_INDEX_CSV = RESULT1_ROOT / "04_3_export_effective_boundaries" / "Otsu" / "04_3_effective_boundary_index.csv"
LAKE_SHAPEFILE = RESULT1_ROOT / "02_Lake_Extraction" / "04_Filtered_Lakes" / "Water_Max_Filtered_Polygons.shp"

SWOT_AUDIT_CSV = RESULT2_ROOT / "06_2_extract_raw_wse" / "Otsu" / "06_1_day_audit_table.csv"
SWOT_RAW_CSV = RESULT2_ROOT / "06_2_extract_raw_wse" / "Otsu" / "06_1_raw_wse_table.csv"
SWOT_BRANCH_DIR = RESULT2_ROOT / "06_3_plot_branchwise_red_chain_with_merge_anchor_redmin5" / "Otsu" / "original" / "06_3_BranchTables" / "vD"
SWOT_NO_NOISE_DIR = RESULT2_ROOT / "06_3_plot_branchwise_red_chain_with_merge_anchor_redmin5" / "Otsu" / "original" / "vd" / "no_noise"
SWOT_PIXC_GDB_DIR = RESULT2_ROOT / "06_1_extract_swot_pixc_to_gdb" / "Otsu" / "06_1_swot_pixc_gdb_by_lake"
SWOT_PIXC_NC_DIR = Path(r"G:\SWOT\Data\PIXC")

_SWOT_PIXC_NC_INDEX: dict[str, Path] | None = None


def ensure_dir(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)


def clean_csv_outputs() -> None:
    if not OUTPUT_DIR.exists():
        return
    for csv_path in OUTPUT_DIR.glob("*.csv"):
        try:
            csv_path.unlink()
        except FileNotFoundError:
            pass


def _read_csv(path: Path) -> pd.DataFrame:
    return pd.read_csv(path, dtype={"lake_id": str}, encoding="utf-8-sig")


def _to_date_str(series: pd.Series) -> pd.Series:
    parsed = pd.to_datetime(series, errors="coerce")
    return parsed.dt.strftime("%Y%m%d")


def _to_date_dt(series: pd.Series) -> pd.Series:
    return pd.to_datetime(series, errors="coerce").dt.normalize()


def _ensure_lake_id_str(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    if "lake_id" in out.columns:
        out["lake_id"] = out["lake_id"].astype(str).str.strip()
    return out


def load_selected_lake_ids() -> list[str]:
    if not SWOT_NO_NOISE_DIR.exists():
        return []
    lake_ids: list[str] = []
    for png_path in sorted(SWOT_NO_NOISE_DIR.glob("*.png")):
        match = re.search(r"Lake_(\d+)_branch_", png_path.name)
        if match:
            lake_ids.append(match.group(1))
    # Keep order stable and unique.
    return list(dict.fromkeys(lake_ids))


def load_lake_metadata(selected_lake_ids: Iterable[str]) -> pd.DataFrame:
    selected_set = {str(v) for v in selected_lake_ids}
    frames: list[pd.DataFrame] = []

    if LAKE_SHAPEFILE.exists():
        try:
            gdf = gpd.read_file(LAKE_SHAPEFILE)
            gdf = pd.DataFrame(gdf.drop(columns="geometry", errors="ignore"))
            gdf = _ensure_lake_id_str(gdf)
            if "lake_id" in gdf.columns:
                frames.append(gdf[gdf["lake_id"].isin(selected_set)].copy())
        except Exception:
            pass

    if not frames and S2_EFFECTIVE_INDEX_CSV.exists():
        idx_df = _ensure_lake_id_str(_read_csv(S2_EFFECTIVE_INDEX_CSV))
        keep_cols = [c for c in idx_df.columns if c in {"lake_id", "size_category", "lake_type"}]
        if keep_cols:
            meta = idx_df[keep_cols].drop_duplicates("lake_id", keep="first").copy()
            frames.append(meta[meta["lake_id"].isin(selected_set)].copy())

    if not frames:
        return pd.DataFrame({"lake_id": list(selected_set)})

    meta_df = pd.concat(frames, ignore_index=True)
    if "lake_id" in meta_df.columns:
        meta_df = meta_df.drop_duplicates("lake_id", keep="first")
    if "size_category" not in meta_df.columns and "lake_size" in meta_df.columns:
        meta_df["size_category"] = meta_df["lake_size"]
    if "lake_size" not in meta_df.columns and "size_category" in meta_df.columns:
        meta_df["lake_size"] = meta_df["size_category"]
    return meta_df


def _load_s2_daily_all() -> pd.DataFrame:
    df = _read_csv(S2_DAILY_ALL_CSV)
    df = _ensure_lake_id_str(df)
    df["date"] = df["date"].astype(str).str.replace(r"\D", "", regex=True)
    df["date_dt"] = pd.to_datetime(df["date"], format="%Y%m%d", errors="coerce")
    return df


def _load_s2_polluted_all() -> pd.DataFrame:
    df = _read_csv(S2_POLLUTED_ALL_CSV)
    df = _ensure_lake_id_str(df)
    df["date"] = df["date"].astype(str).str.replace(r"\D", "", regex=True)
    return df


def _load_s2_effective_index() -> pd.DataFrame:
    if not S2_EFFECTIVE_INDEX_CSV.exists():
        return pd.DataFrame()
    df = _read_csv(S2_EFFECTIVE_INDEX_CSV)
    df = _ensure_lake_id_str(df)
    df["date"] = df["date"].astype(str).str.replace(r"\D", "", regex=True)
    df["water_period_start"] = pd.to_datetime(df["water_period_start"], errors="coerce")
    df["water_period_end"] = pd.to_datetime(df["water_period_end"], errors="coerce")
    return df


def _build_ndsi_index() -> dict[str, Path]:
    index: dict[str, Path] = {}
    if not S2_NDSI_DIR.exists():
        return index
    for tif_path in S2_NDSI_DIR.glob("*.tif"):
        match = re.search(r"_(\d{8}T\d{6})_", tif_path.name)
        if match:
            index[match.group(1)] = tif_path
    return index


def _build_cloudmask_index() -> set[str]:
    out: set[str] = set()
    if not S2_CLOUD_DIR.exists():
        return out
    for tif_path in S2_CLOUD_DIR.glob("*.tif"):
        match = re.search(r"_(\d{8}T\d{6})_", tif_path.name)
        if match:
            out.add(match.group(1))
    return out


def _preview_scene_parts(preview_name: str) -> tuple[str, str, str, str] | None:
    match = re.match(r"overlay_(\d+)_(\d{8})_(\d{6})_([A-Za-z0-9]+)\.png", preview_name)
    if not match:
        return None
    _, date_str, time_str, platform = match.groups()
    scene_key = f"{date_str}_{time_str}_{platform}"
    return date_str, time_str, platform, scene_key


def _water_status_from_fixed_window(date_str: str) -> str:
    date_dt = pd.to_datetime(date_str, format="%Y%m%d", errors="coerce")
    if pd.isna(date_dt):
        return ""
    return "有水" if S2_START <= date_dt <= S2_END else "无水"


def _boundary_status_from_row(row: pd.Series) -> str:
    if row is None or row.empty:
        return ""
    if pd.notna(row.get("effective_features")) and float(row.get("effective_features", 0) or 0) > 0:
        return "有效边界"
    if pd.notna(row.get("jpg_bucket")):
        value = str(row["jpg_bucket"])
        if "有效边界" in value:
            return "有效边界"
        if "无效边界" in value:
            return "非有效边界"
    if pd.notna(row.get("final_class")):
        return "有效边界" if str(row["final_class"]).lower() == "water" else "非有效边界"
    return ""


def _format_scene_time(date_str: str, time_str: str) -> str:
    if len(date_str) != 8 or len(time_str) < 4:
        return ""
    return f"{date_str[:4]}-{date_str[4:6]}-{date_str[6:8]} {time_str[:2]}:{time_str[2:4]}"


def build_s2_status_records(selected_lake_ids: Iterable[str], lake_meta: pd.DataFrame | None = None) -> pd.DataFrame:
    selected_set = {str(v) for v in selected_lake_ids}
    daily_df = _load_s2_daily_all()
    daily_df = daily_df[daily_df["lake_id"].isin(selected_set)].copy()

    polluted_df = _load_s2_polluted_all()
    polluted_df = polluted_df[polluted_df["lake_id"].isin(selected_set)].copy()
    polluted_stats = (
        polluted_df.groupby(["lake_id", "date"], as_index=False)
        .agg(
            polluted_image_count=("source_image", "count"),
            polluted_cloud_ratio_mean=("cloud_ratio", "mean"),
            polluted_cloud_ratio_max=("cloud_ratio", "max"),
        )
        if not polluted_df.empty
        else pd.DataFrame(columns=["lake_id", "date", "polluted_image_count", "polluted_cloud_ratio_mean", "polluted_cloud_ratio_max"])
    )

    effective_df = _load_s2_effective_index()
    effective_df = effective_df[effective_df["lake_id"].isin(selected_set)].copy() if not effective_df.empty else effective_df
    if not effective_df.empty:
        effective_stats = (
            effective_df.groupby(["lake_id", "date"], as_index=False)
            .agg(
                effective_features=("effective_features", "sum"),
                invalid_features=("invalid_features", "sum"),
                total_features=("total_features", "sum"),
                boundary_rows=("preview_name", "count"),
                effective_water_rows=("final_class", lambda s: int((s.astype(str).str.lower() == "water").sum())),
            )
        )
    else:
        effective_stats = pd.DataFrame(
            columns=["lake_id", "date", "effective_features", "invalid_features", "total_features", "boundary_rows", "effective_water_rows"]
        )

    out = daily_df.merge(polluted_stats, on=["lake_id", "date"], how="left")
    out = out.merge(effective_stats, on=["lake_id", "date"], how="left")

    if lake_meta is not None and not lake_meta.empty:
        meta_cols = [c for c in lake_meta.columns if c in {"lake_id", "lake_size", "size_category", "lake_type"}]
        if meta_cols:
            out = out.merge(lake_meta[meta_cols].drop_duplicates("lake_id", keep="first"), on="lake_id", how="left")

    if "lake_size" not in out.columns and "size_category" in out.columns:
        out["lake_size"] = out["size_category"]
    if "size_category" not in out.columns and "lake_size" in out.columns:
        out["size_category"] = out["lake_size"]

    out["source"] = "S2"
    out["date_dt"] = pd.to_datetime(out["date"], format="%Y%m%d", errors="coerce")
    out["in_fixed_window"] = out["date_dt"].between(S2_START, S2_END, inclusive="both")
    out["cloud_pollution_rate"] = np.where(
        out["valid_pixel_images"].fillna(0).astype(float) > 0,
        out["cloud_polluted_images"].fillna(0).astype(float) / out["valid_pixel_images"].fillna(0).astype(float),
        np.nan,
    )
    out["valid_pixel_rate"] = np.where(
        out["total_images"].fillna(0).astype(float) > 0,
        out["valid_pixel_images"].fillna(0).astype(float) / out["total_images"].fillna(0).astype(float),
        np.nan,
    )
    out = out.sort_values(["lake_id", "date"]).reset_index(drop=True)
    return out


def build_s2_reference_status_records(selected_lake_ids: Iterable[str], lake_meta: pd.DataFrame | None = None) -> pd.DataFrame:
    selected_ids = [str(v) for v in selected_lake_ids]
    meta_lookup = {}
    if lake_meta is not None and not lake_meta.empty:
        meta_df = lake_meta.copy()
        meta_df["lake_id"] = meta_df["lake_id"].astype(str)
        meta_lookup = meta_df.set_index("lake_id").to_dict("index")

    polluted_df = _load_s2_polluted_all()
    polluted_df["preview_image"] = polluted_df["preview_image"].astype(str)
    polluted_set = set(polluted_df["preview_image"].tolist())

    effective_df = _load_s2_effective_index()
    effective_lookup: dict[tuple[str, str], pd.Series] = {}
    if not effective_df.empty:
        for _, row in effective_df.iterrows():
            effective_lookup[(str(row["lake_id"]), str(row["preview_name"]))] = row

    ndsi_index = _build_ndsi_index()
    cloud_index = _build_cloudmask_index()

    rows: list[dict[str, object]] = []
    for lake_id in selected_ids:
        lake_meta_row = meta_lookup.get(lake_id, {})
        lake_type = str(lake_meta_row.get("lake_type") or lake_meta_row.get("size_category") or lake_meta_row.get("lake_size") or "")
        base_dir = S2_PREVIEW_ROOT / lake_type / lake_id / "01_Original_Base"
        clean_valid_dir = S2_PREVIEW_ROOT / lake_type / lake_id / "02_Clean_Valid"
        if not base_dir.exists():
            continue

        clean_valid_names = {p.name for p in clean_valid_dir.glob("*.png")} if clean_valid_dir.exists() else set()
        for preview_path in sorted(base_dir.glob("*.png")):
            parts = _preview_scene_parts(preview_path.name)
            if parts is None:
                continue
            date_str, time_str, platform, scene_key = parts
            dt_key = f"{date_str}T{time_str}"
            image_path = ndsi_index.get(dt_key)
            image_status = "cloud_covered" if preview_path.name in polluted_set else ("normal" if preview_path.name in clean_valid_names else "invalid")
            effective_row = effective_lookup.get((lake_id, preview_path.name))
            rows.append(
                {
                    "湖泊编号": int(lake_id),
                    "湖泊面积_平方米": lake_meta_row.get("area"),
                    "影像文件名": image_path.name if image_path else "",
                    "影像日期": date_str,
                    "影像时间": _format_scene_time(date_str, time_str),
                    "场景键": scene_key,
                    "云掩膜是否存在": dt_key in cloud_index,
                    "水期状态": _water_status_from_fixed_window(date_str),
                    "边界状态": _boundary_status_from_row(effective_row) or "",
                    "影像状态": image_status,
                    "影像路径": str(image_path) if image_path else "",
                }
            )

    out = pd.DataFrame(rows)
    if out.empty:
        return out
    out = out.sort_values(["湖泊编号", "影像日期", "影像时间", "场景键"]).reset_index(drop=True)
    return out


def build_s2_reference_daily_records(status_df: pd.DataFrame) -> pd.DataFrame:
    if status_df.empty:
        return pd.DataFrame(
            columns=[
                "湖泊编号",
                "湖泊面积_平方米",
                "日期",
                "日统计状态",
                "采用水期状态",
                "采用边界状态",
                "采用影像文件名",
                "采用影像时间",
                "采用卫星平台",
                "采用场景键",
                "采用影像路径",
                "当日影像次数",
                "当日正常次数",
                "当日云遮挡次数",
                "当日无效次数",
            ]
        )

    work = status_df.copy()
    priority = {"normal": 0, "cloud_covered": 1, "invalid": 2}
    work["_priority"] = work["影像状态"].map(priority).fillna(9)
    work = work.sort_values(["湖泊编号", "影像日期", "_priority", "影像时间", "场景键"])

    rows: list[dict[str, object]] = []
    for (lake_id, date_str), group in work.groupby(["湖泊编号", "影像日期"], dropna=False):
        adopted = group.iloc[0]
        status_set = set(group["影像状态"].astype(str))
        if "normal" in status_set:
            day_status = "normal"
        elif "cloud_covered" in status_set:
            day_status = "cloud_covered"
        else:
            day_status = "invalid"
        rows.append(
            {
                "湖泊编号": lake_id,
                "湖泊面积_平方米": adopted.get("湖泊面积_平方米"),
                "日期": date_str,
                "日统计状态": day_status,
                "采用水期状态": adopted.get("水期状态", ""),
                "采用边界状态": adopted.get("边界状态", ""),
                "采用影像文件名": adopted.get("影像文件名", ""),
                "采用影像时间": adopted.get("影像时间", ""),
                "采用场景键": adopted.get("场景键", ""),
                "采用影像路径": adopted.get("影像路径", ""),
                "当日影像次数": len(group),
                "当日正常次数": int((group["影像状态"] == "normal").sum()),
                "当日云遮挡次数": int((group["影像状态"] == "cloud_covered").sum()),
                "当日无效次数": int((group["影像状态"] == "invalid").sum()),
            }
        )
    return pd.DataFrame(rows).sort_values(["湖泊编号", "日期"]).reset_index(drop=True)


def _pick_best_raw_row(raw_group: pd.DataFrame) -> pd.Series:
    if raw_group.empty:
        return pd.Series(dtype=object)
    if "wse_stage" in raw_group.columns:
        preferred = raw_group[raw_group["wse_stage"].astype(str).str.lower() == "post_sigma"]
        if not preferred.empty:
            return preferred.iloc[0]
    return raw_group.iloc[0]


def _load_branch_noise_table(selected_lake_ids: Iterable[str]) -> pd.DataFrame:
    selected_set = {str(v) for v in selected_lake_ids}
    if not SWOT_BRANCH_DIR.exists():
        return pd.DataFrame(columns=["lake_id", "branch_index", "date", "is_noise", "noise_reason"])

    frames: list[pd.DataFrame] = []
    for csv_path in sorted(SWOT_BRANCH_DIR.glob("lake_*_branch_*_stage_vD.csv")):
        match = re.match(r"lake_(\d+)_branch_(\d+)_stage_vD\.csv", csv_path.name)
        if not match:
            continue
        lake_id, branch_index = match.groups()
        if lake_id not in selected_set:
            continue
        try:
            df = pd.read_csv(csv_path)
        except Exception:
            continue
        if df.empty:
            continue
        if "date" not in df.columns:
            continue
        df = df.copy()
        df["lake_id"] = lake_id
        df["branch_index"] = pd.to_numeric(branch_index, errors="coerce")
        if "date" in df.columns:
            df["date"] = df["date"].astype(str).str.replace(r"\D", "", regex=True)
        if "is_noise" not in df.columns:
            df["is_noise"] = False
        if "noise_reason" not in df.columns:
            df["noise_reason"] = ""
        frames.append(df[["lake_id", "branch_index", "date", "is_noise", "noise_reason"]].copy())

    if not frames:
        return pd.DataFrame(columns=["lake_id", "branch_index", "date", "is_noise", "noise_reason"])

    noise_df = pd.concat(frames, ignore_index=True)
    noise_df["is_noise"] = noise_df["is_noise"].astype(bool)
    noise_df = (
        noise_df.groupby(["lake_id", "branch_index", "date"], as_index=False)
        .agg(
            is_noise=("is_noise", lambda s: bool(s.all())),
            noise_reason=("noise_reason", lambda s: "|".join(sorted({str(v) for v in s if pd.notna(v) and str(v)}))),
        )
    )
    return noise_df


def load_raw_wse_table() -> pd.DataFrame:
    df = pd.read_csv(SWOT_RAW_CSV)
    df = _ensure_lake_id_str(df)
    if "date" in df.columns:
        df["date"] = df["date"].astype(str).str.replace(r"\D", "", regex=True)
    if "date_dt" in df.columns:
        df["date_dt"] = pd.to_datetime(df["date_dt"], errors="coerce")
    else:
        df["date_dt"] = pd.to_datetime(df["date"], format="%Y%m%d", errors="coerce")
    if "branch_index" in df.columns:
        df["branch_index"] = pd.to_numeric(df["branch_index"], errors="coerce")
    return df


def load_day_audit_table() -> pd.DataFrame:
    df = pd.read_csv(SWOT_AUDIT_CSV)
    df = _ensure_lake_id_str(df)
    if "date" in df.columns:
        df["date"] = df["date"].astype(str).str.replace(r"\D", "", regex=True)
    if "branch_index" in df.columns:
        df["branch_index"] = pd.to_numeric(df["branch_index"], errors="coerce")
    return df


def _build_swot_pixc_nc_index() -> dict[str, Path]:
    global _SWOT_PIXC_NC_INDEX
    if _SWOT_PIXC_NC_INDEX is not None:
        return _SWOT_PIXC_NC_INDEX

    index: dict[str, Path] = {}
    if SWOT_PIXC_NC_DIR.exists():
        pattern = re.compile(
            r"SWOT_L2_HR_PIXC_(\d+)_(\d+)_(\d+[LR])_(\d{8}T\d{6})_(\d{8}T\d{6})_([A-Z0-9]+)_01\.nc$",
            re.IGNORECASE,
        )
        for nc_path in SWOT_PIXC_NC_DIR.rglob("*.nc"):
            match = pattern.match(nc_path.name)
            if not match:
                continue
            _, pass_id, tile_id, start_dt, _, product = match.groups()
            key = f"pixc_{start_dt}_{pass_id}_{tile_id}_{product}_01".lower()
            index.setdefault(key, nc_path)

    _SWOT_PIXC_NC_INDEX = index
    return index


def _resolve_swot_granule(source_layers: object) -> tuple[str, bool, str]:
    if pd.isna(source_layers):
        return "", False, ""

    nc_index = _build_swot_pixc_nc_index()
    tokens = [str(v).strip() for v in str(source_layers).split("|") if str(v).strip()]
    seen: set[str] = set()
    for token in tokens:
        key = token.lower()
        if key in seen:
            continue
        seen.add(key)
        nc_path = nc_index.get(key)
        if nc_path is not None:
            return nc_path.name, True, str(nc_path)
    return "", False, ""


def _parse_swot_granule_tokens(source_layers: object, reason_detail: object) -> list[str]:
    pattern = re.compile(r"pixc_\d{8}T\d{6}_\d+_\d+[LR]_[A-Z0-9]+_01", re.IGNORECASE)
    tokens: list[str] = []
    seen: set[str] = set()

    def _push(text: object) -> None:
        if pd.isna(text):
            return
        for token in pattern.findall(str(text)):
            key = token.lower()
            if key in seen:
                continue
            seen.add(key)
            tokens.append(token)

    _push(source_layers)
    _push(reason_detail)
    return tokens


def _resolve_swot_granule_token(token: str) -> tuple[str, bool, str]:
    if not token:
        return "", False, ""
    nc_index = _build_swot_pixc_nc_index()
    nc_path = nc_index.get(str(token).strip().lower())
    if nc_path is None:
        return token, False, ""
    return nc_path.name, True, str(nc_path)


def build_swot_status_records(selected_lake_ids: Iterable[str], lake_meta: pd.DataFrame | None = None) -> pd.DataFrame:
    selected_set = {str(v) for v in selected_lake_ids}
    audit_df = load_day_audit_table()
    audit_df = audit_df[audit_df["lake_id"].isin(selected_set)].copy()
    raw_df = load_raw_wse_table()
    raw_df = raw_df[raw_df["lake_id"].isin(selected_set)].copy()

    noise_df = _load_branch_noise_table(selected_set)
    if not noise_df.empty:
        noise_df["branch_index"] = pd.to_numeric(noise_df["branch_index"], errors="coerce")

    best_raw_rows: list[pd.Series] = []
    if not raw_df.empty:
        grouped = raw_df.groupby(["lake_id", "date", "branch_index"], dropna=False)
        for _, group in grouped:
            best_raw_rows.append(_pick_best_raw_row(group))
    raw_best_df = pd.DataFrame(best_raw_rows) if best_raw_rows else pd.DataFrame()

    out = audit_df.copy()
    if not raw_best_df.empty:
        merge_cols = [c for c in ["lake_id", "date", "branch_index"] if c in out.columns and c in raw_best_df.columns]
        if merge_cols:
            raw_best_df = raw_best_df.drop_duplicates(merge_cols, keep="first")
            out = out.merge(raw_best_df, on=merge_cols, how="left", suffixes=("", "_raw"))

    if not noise_df.empty and {"lake_id", "date", "branch_index"}.issubset(out.columns):
        out = out.merge(noise_df, on=["lake_id", "date", "branch_index"], how="left", suffixes=("", "_noise"))

    if lake_meta is not None and not lake_meta.empty:
        meta_cols = [c for c in lake_meta.columns if c in {"lake_id", "lake_size", "size_category", "lake_type"}]
        if meta_cols:
            out = out.merge(lake_meta[meta_cols].drop_duplicates("lake_id", keep="first"), on="lake_id", how="left")

    if "lake_size" not in out.columns and "size_category" in out.columns:
        out["lake_size"] = out["size_category"]
    if "size_category" not in out.columns and "lake_size" in out.columns:
        out["size_category"] = out["lake_size"]

    out["source"] = "SWOT"
    out["date_dt"] = pd.to_datetime(out["date"], format="%Y%m%d", errors="coerce")
    out["in_fixed_window"] = out["date_dt"].between(S2_START, S2_END, inclusive="both")
    if "is_noise" not in out.columns:
        out["is_noise"] = False
    out["is_noise"] = out["is_noise"].fillna(False).astype(bool)
    if "noise_reason" not in out.columns:
        out["noise_reason"] = ""
    out["noise_reason"] = out["noise_reason"].fillna("")
    out["has_valid_wse"] = out.get("status", pd.Series(dtype=object)).astype(str).eq("valid_wse")
    out = out.sort_values(["lake_id", "date", "branch_index"], na_position="last").reset_index(drop=True)
    return out


def _map_swot_reason(status: str, reason_short: str) -> str:
    status = str(status)
    reason_short = str(reason_short)
    if status == "no_pixc":
        return "无PIXC"
    if reason_short == "BaseFilter":
        return "基础筛选后无有效点"
    if reason_short == "ClipEmpty":
        return "裁剪后无有效点"
    return ""


def build_swot_reference_status_records(selected_lake_ids: Iterable[str], lake_meta: pd.DataFrame | None = None) -> pd.DataFrame:
    raw_df = build_swot_status_records(selected_lake_ids, lake_meta)
    if raw_df.empty:
        return pd.DataFrame()

    meta_lookup = {}
    if lake_meta is not None and not lake_meta.empty:
        meta_df = lake_meta.copy()
        meta_df["lake_id"] = meta_df["lake_id"].astype(str)
        meta_lookup = meta_df.set_index("lake_id").to_dict("index")

    rows: list[dict[str, object]] = []
    for _, row in raw_df.iterrows():
        lake_id = str(row.get("lake_id"))
        lake_gdb = SWOT_PIXC_GDB_DIR / f"swot_pixc_{lake_id}.gdb"
        time_value = ""
        if pd.notna(row.get("date_dt")):
            dt = pd.to_datetime(row["date_dt"], errors="coerce")
            if pd.notna(dt):
                time_value = dt.strftime("%Y-%m-%d")
        reason = _map_swot_reason(row.get("status", ""), row.get("reason_short", ""))
        noise_label = ""
        if str(row.get("status")) == "valid_wse":
            noise_label = "噪声" if bool(row.get("is_noise", False)) else "非噪声"

        rows.append(
            {
                "湖泊编号": int(lake_id),
                "分支编号": row.get("branch_index"),
                "分支ID": row.get("branch_id", ""),
                "PIXC版本": row.get("pixc_version", ""),
                "SWOT日期": row.get("date", ""),
                "SWOT时间": time_value,
                "审核状态": row.get("status", ""),
                "原因": reason,
                "是否噪声": noise_label,
                "当前匹配图层": row.get("source_layers", ""),
                "水期状态": _water_status_from_fixed_window(str(row.get("date", ""))),
                "GDB图层是否存在": lake_gdb.exists(),
                "GDB路径": str(lake_gdb),
            }
        )
    out = pd.DataFrame(rows)
    return out.sort_values(["湖泊编号", "SWOT日期", "分支编号"]).reset_index(drop=True)


def build_swot_reference_daily_records(status_df: pd.DataFrame) -> pd.DataFrame:
    if status_df.empty:
        return pd.DataFrame(
            columns=[
                "湖泊编号",
                "统计层级",
                "统计对象",
                "日期",
                "日统计状态",
                "采用水期状态",
                "采用SWOT时间",
                "采用分支编号",
                "采用分支ID",
                "采用是否噪声",
                "采用匹配图层",
                "采用原因",
                "采用记录GDB图层是否存在",
                "采用记录GDB路径",
                "当日记录次数",
                "当日有效次数",
                "当日无效次数",
            ]
        )

    work = status_df.copy()
    work["_valid_rank"] = np.where(work["审核状态"].astype(str) == "valid_wse", 0, 1)
    work["_noise_rank"] = np.where(work["是否噪声"].astype(str) == "非噪声", 0, 1)
    work = work.sort_values(["湖泊编号", "SWOT日期", "_valid_rank", "_noise_rank", "SWOT时间"])

    rows: list[dict[str, object]] = []
    for (lake_id, date_str), group in work.groupby(["湖泊编号", "SWOT日期"], dropna=False):
        adopted = group.iloc[0]
        valid_mask = group["审核状态"].astype(str) == "valid_wse"
        day_status = "valid" if valid_mask.any() else "invalid"
        branch_no = adopted.get("分支编号")
        rows.append(
            {
                "湖泊编号": lake_id,
                "统计层级": "branch",
                "统计对象": f"L{int(lake_id)}_B{int(branch_no)}" if pd.notna(branch_no) else f"L{int(lake_id)}",
                "日期": date_str,
                "日统计状态": day_status,
                "采用水期状态": adopted.get("水期状态", ""),
                "采用SWOT时间": adopted.get("SWOT时间", ""),
                "采用分支编号": branch_no,
                "采用分支ID": adopted.get("分支ID", ""),
                "采用是否噪声": adopted.get("是否噪声", ""),
                "采用匹配图层": adopted.get("当前匹配图层", ""),
                "采用原因": adopted.get("原因", ""),
                "采用记录GDB图层是否存在": adopted.get("GDB图层是否存在", False),
                "采用记录GDB路径": adopted.get("GDB路径", ""),
                "当日记录次数": len(group),
                "当日有效次数": int(valid_mask.sum()),
                "当日无效次数": int((~valid_mask).sum()),
            }
        )

    return pd.DataFrame(rows).sort_values(["湖泊编号", "日期"]).reset_index(drop=True)


def build_daily_summary(status_df: pd.DataFrame) -> pd.DataFrame:
    if status_df.empty:
        return pd.DataFrame()
    work = status_df.copy()
    if "date_dt" not in work.columns and "date" in work.columns:
        work["date_dt"] = pd.to_datetime(work["date"], format="%Y%m%d", errors="coerce")
    group_cols = ["date"]
    agg: dict[str, tuple[str, str]] = {}
    for col in work.columns:
        if col in {"lake_id", "date", "date_dt"}:
            continue
        if pd.api.types.is_bool_dtype(work[col]):
            agg[col] = (col, "sum")
        elif pd.api.types.is_numeric_dtype(work[col]):
            agg[col] = (col, "sum")
    out = work.groupby(group_cols, as_index=False).agg(**agg) if agg else work[group_cols].drop_duplicates().copy()
    if "date_dt" not in out.columns:
        out["date_dt"] = pd.to_datetime(out["date"], format="%Y%m%d", errors="coerce")
    if "lake_id" in work.columns:
        lake_counts = work.groupby("date", as_index=False)["lake_id"].nunique().rename(columns={"lake_id": "lake_count"})
        out = out.merge(lake_counts, on="date", how="left")
    out["in_fixed_window"] = out["date_dt"].between(S2_START, S2_END, inclusive="both")
    out = out.sort_values("date").reset_index(drop=True)
    return out


def build_lake_summary(status_df: pd.DataFrame, only_fixed_window: bool = False) -> pd.DataFrame:
    if status_df.empty:
        return pd.DataFrame()
    work = status_df.copy()
    if only_fixed_window and "in_fixed_window" in work.columns:
        work = work[work["in_fixed_window"]].copy()
    if work.empty:
        return pd.DataFrame()

    group_cols = ["lake_id"]
    if "lake_size" in work.columns:
        group_cols.append("lake_size")
    if "size_category" in work.columns and "size_category" not in group_cols:
        group_cols.append("size_category")

    agg_map: dict[str, tuple[str, str]] = {
        "date": ("date", "nunique"),
    }
    if "status" in work.columns:
        for status_name in ["valid_wse", "no_valid_wse", "no_pixc"]:
            col_name = f"{status_name}_count"
            agg_map[col_name] = ("status", lambda s, name=status_name: int((s.astype(str) == name).sum()))
    if "is_noise" in work.columns:
        agg_map["noise_count"] = ("is_noise", "sum")
    if "has_valid_wse" in work.columns:
        agg_map["valid_wse_days"] = ("has_valid_wse", "sum")

    out = work.groupby(group_cols, as_index=False).agg(**agg_map)
    out = out.rename(columns={"date": "observation_days"})
    out["in_fixed_window"] = only_fixed_window
    out = out.sort_values(group_cols).reset_index(drop=True)
    return out


def build_date_summary(status_df: pd.DataFrame, only_fixed_window: bool = False) -> pd.DataFrame:
    if status_df.empty:
        return pd.DataFrame()
    work = status_df.copy()
    if only_fixed_window and "in_fixed_window" in work.columns:
        work = work[work["in_fixed_window"]].copy()
    if work.empty:
        return pd.DataFrame()

    agg_map: dict[str, tuple[str, str]] = {
        "lake_id": ("lake_id", "nunique"),
    }
    if "status" in work.columns:
        for status_name in ["valid_wse", "no_valid_wse", "no_pixc"]:
            agg_map[f"{status_name}_count"] = ("status", lambda s, name=status_name: int((s.astype(str) == name).sum()))
    if "is_noise" in work.columns:
        agg_map["noise_count"] = ("is_noise", "sum")
    if "has_valid_wse" in work.columns:
        agg_map["valid_wse_days"] = ("has_valid_wse", "sum")

    out = work.groupby("date", as_index=False).agg(**agg_map)
    out = out.rename(columns={"lake_id": "lake_count"})
    out["date_dt"] = pd.to_datetime(out["date"], format="%Y%m%d", errors="coerce")
    out["in_fixed_window"] = only_fixed_window
    out = out.sort_values("date").reset_index(drop=True)
    return out


def build_date_overlap_summary(s2_df: pd.DataFrame, swot_df: pd.DataFrame) -> pd.DataFrame:
    if s2_df.empty and swot_df.empty:
        return pd.DataFrame()

    s2_work = s2_df[["lake_id", "date", "in_fixed_window"]].drop_duplicates().copy() if not s2_df.empty else pd.DataFrame(columns=["lake_id", "date", "in_fixed_window"])
    swot_work = swot_df[["lake_id", "date", "in_fixed_window"]].drop_duplicates().copy() if not swot_df.empty else pd.DataFrame(columns=["lake_id", "date", "in_fixed_window"])
    lake_ids = sorted(set(s2_work["lake_id"]).union(set(swot_work["lake_id"])))

    rows: list[dict[str, object]] = []
    for lake_id in lake_ids:
        s2_all = set(s2_work.loc[s2_work["lake_id"] == lake_id, "date"].astype(str))
        swot_all = set(swot_work.loc[swot_work["lake_id"] == lake_id, "date"].astype(str))
        s2_fix = set(s2_work.loc[(s2_work["lake_id"] == lake_id) & (s2_work["in_fixed_window"]), "date"].astype(str))
        swot_fix = set(swot_work.loc[(swot_work["lake_id"] == lake_id) & (swot_work["in_fixed_window"]), "date"].astype(str))

        rows.append(
            {
                "lake_id": lake_id,
                "s2_total_dates": len(s2_all),
                "swot_total_dates": len(swot_all),
                "intersection_dates": len(s2_all & swot_all),
                "union_dates": len(s2_all | swot_all),
                "s2_only_dates": len(s2_all - swot_all),
                "swot_only_dates": len(swot_all - s2_all),
                "s2_fixed_window_dates": len(s2_fix),
                "swot_fixed_window_dates": len(swot_fix),
                "intersection_fixed_window_dates": len(s2_fix & swot_fix),
                "union_fixed_window_dates": len(s2_fix | swot_fix),
                "s2_only_fixed_window_dates": len(s2_fix - swot_fix),
                "swot_only_fixed_window_dates": len(swot_fix - s2_fix),
            }
        )

    return pd.DataFrame(rows).sort_values("lake_id").reset_index(drop=True)


def build_resummary(sheet_df: pd.DataFrame) -> pd.DataFrame:
    if sheet_df.empty:
        return sheet_df.copy()
    out = sheet_df.copy()
    if "date" in out.columns:
        out["date"] = out["date"].astype(str)
    for col in out.columns:
        if pd.api.types.is_float_dtype(out[col]):
            out[col] = out[col].round(6)
    return out.reset_index(drop=True)


def _set_sheet_style(ws) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells[:200]:
            value = cell.value
            if value is None:
                continue
            max_len = max(max_len, len(str(value)))
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 40)


def format_excel_workbook(path: Path) -> None:
    if not path.exists():
        return
    wb = load_workbook(path)
    for ws in wb.worksheets:
        _set_sheet_style(ws)
    wb.save(path)


def build_swot_reference_status_records(selected_lake_ids: Iterable[str], lake_meta: pd.DataFrame | None = None) -> pd.DataFrame:
    raw_df = build_swot_status_records(selected_lake_ids, lake_meta)
    if raw_df.empty:
        return pd.DataFrame()

    rows: list[dict[str, object]] = []
    for _, row in raw_df.iterrows():
        lake_id = str(row.get("lake_id"))
        lake_gdb = SWOT_PIXC_GDB_DIR / f"swot_pixc_{lake_id}.gdb"
        granule_tokens = _parse_swot_granule_tokens(row.get("source_layers", ""), row.get("reason_detail", ""))
        if not granule_tokens:
            granule_tokens = [""]

        time_value = ""
        if pd.notna(row.get("date_dt")):
            dt = pd.to_datetime(row["date_dt"], errors="coerce")
            if pd.notna(dt):
                time_value = dt.strftime("%Y-%m-%d")

        reason = _map_swot_reason(row.get("status", ""), row.get("reason_short", ""))
        noise_label = ""
        if str(row.get("status")) == "valid_wse":
            noise_label = "\u566a\u58f0" if bool(row.get("is_noise", False)) else "\u975e\u566a\u58f0"

        for granule_token in granule_tokens:
            granule_name, nc_exists, nc_path = _resolve_swot_granule_token(granule_token)
            gdb_exists = bool(nc_exists and lake_gdb.exists())
            gdb_path = str(lake_gdb) if gdb_exists else ""

            rows.append(
                {
                    "\u6e56\u6cca\u7f16\u53f7": int(lake_id),
                    "\u5206\u652f\u7f16\u53f7": row.get("branch_index"),
                    "\u5206\u652fID": row.get("branch_id", ""),
                    "PIXC\u7248\u672c": row.get("pixc_version", ""),
                    "SWOT\u65e5\u671f": row.get("date", ""),
                    "SWOT\u65f6\u95f4": time_value,
                    "\u5ba1\u6838\u72b6\u6001": row.get("status", ""),
                    "\u539f\u56e0": reason,
                    "\u662f\u5426\u566a\u58f0": noise_label,
                    "\u5f53\u524d\u5339\u914d\u56fe\u5c42": granule_token,
                    "Granule\u540d\u79f0": granule_name,
                    "\u6c34\u671f\u72b6\u6001": _water_status_from_fixed_window(str(row.get("date", ""))),
                    "NC\u6587\u4ef6\u662f\u5426\u5b58\u5728": nc_exists,
                    "NC\u6587\u4ef6\u8def\u5f84": nc_path,
                    "GDB\u56fe\u5c42\u662f\u5426\u5b58\u5728": gdb_exists,
                    "GDB\u8def\u5f84": gdb_path,
                }
            )

    out = pd.DataFrame(rows)
    return out.sort_values(["\u6e56\u6cca\u7f16\u53f7", "SWOT\u65e5\u671f", "\u5206\u652f\u7f16\u53f7"]).reset_index(drop=True)


def build_swot_reference_daily_records(status_df: pd.DataFrame) -> pd.DataFrame:
    if status_df.empty:
        return pd.DataFrame(
            columns=[
                "\u6e56\u6cca\u7f16\u53f7",
                "\u7edf\u8ba1\u5c42\u7ea7",
                "\u7edf\u8ba1\u5bf9\u8c61",
                "\u65e5\u671f",
                "\u65e5\u7edf\u8ba1\u72b6\u6001",
                "\u91c7\u7528\u6c34\u671f\u72b6\u6001",
                "\u91c7\u7528SWOT\u65f6\u95f4",
                "\u91c7\u7528\u5206\u652f\u7f16\u53f7",
                "\u91c7\u7528\u5206\u652fID",
                "\u91c7\u7528\u5339\u914d\u56fe\u5c42",
                "\u91c7\u7528Granule\u540d\u79f0",
                "\u91c7\u7528\u539f\u56e0",
                "\u91c7\u7528\u8bb0\u5f55NC\u6587\u4ef6\u662f\u5426\u5b58\u5728",
                "\u91c7\u7528\u8bb0\u5f55NC\u6587\u4ef6\u8def\u5f84",
                "\u91c7\u7528\u8bb0\u5f55GDB\u56fe\u5c42\u662f\u5426\u5b58\u5728",
                "\u91c7\u7528\u8bb0\u5f55GDB\u8def\u5f84",
                "\u5f53\u65e5\u8bb0\u5f55\u6b21\u6570",
                "\u5f53\u65e5\u6709\u6548\u6b21\u6570",
                "\u5f53\u65e5\u65e0\u6548\u6b21\u6570",
            ]
        )

    c_lake = "\u6e56\u6cca\u7f16\u53f7"
    c_date = "SWOT\u65e5\u671f"
    c_status = "\u5ba1\u6838\u72b6\u6001"
    c_noise = "\u662f\u5426\u566a\u58f0"
    c_branch_no = "\u5206\u652f\u7f16\u53f7"
    c_branch_id = "\u5206\u652fID"

    work = status_df.copy()
    work["_valid_rank"] = np.where(work[c_status].astype(str) == "valid_wse", 0, 1)
    work["_noise_rank"] = np.where(work[c_noise].astype(str) == "\u975e\u566a\u58f0", 0, 1)
    work = work.sort_values([c_lake, c_date, "_valid_rank", "_noise_rank", "SWOT\u65f6\u95f4"])

    rows: list[dict[str, object]] = []
    for (lake_id, date_str), group in work.groupby([c_lake, c_date], dropna=False):
        adopted = group.iloc[0]
        valid_mask = group[c_status].astype(str) == "valid_wse"
        day_status = "valid" if valid_mask.any() else "invalid"
        branch_no = adopted.get(c_branch_no)
        rows.append(
            {
                "\u6e56\u6cca\u7f16\u53f7": lake_id,
                "\u7edf\u8ba1\u5c42\u7ea7": "branch",
                "\u7edf\u8ba1\u5bf9\u8c61": f"L{int(lake_id)}_B{int(branch_no)}" if pd.notna(branch_no) else f"L{int(lake_id)}",
                "\u65e5\u671f": date_str,
                "\u65e5\u7edf\u8ba1\u72b6\u6001": day_status,
                "\u91c7\u7528\u6c34\u671f\u72b6\u6001": adopted.get("\u6c34\u671f\u72b6\u6001", ""),
                "\u91c7\u7528SWOT\u65f6\u95f4": adopted.get("SWOT\u65f6\u95f4", ""),
                "\u91c7\u7528\u5206\u652f\u7f16\u53f7": branch_no,
                "\u91c7\u7528\u5206\u652fID": adopted.get(c_branch_id, ""),
                "\u91c7\u7528\u5339\u914d\u56fe\u5c42": adopted.get("\u5f53\u524d\u5339\u914d\u56fe\u5c42", ""),
                "\u91c7\u7528Granule\u540d\u79f0": adopted.get("Granule\u540d\u79f0", ""),
                "\u91c7\u7528\u539f\u56e0": adopted.get("\u539f\u56e0", ""),
                "\u91c7\u7528\u8bb0\u5f55NC\u6587\u4ef6\u662f\u5426\u5b58\u5728": adopted.get("NC\u6587\u4ef6\u662f\u5426\u5b58\u5728", False),
                "\u91c7\u7528\u8bb0\u5f55NC\u6587\u4ef6\u8def\u5f84": adopted.get("NC\u6587\u4ef6\u8def\u5f84", ""),
                "\u91c7\u7528\u8bb0\u5f55GDB\u56fe\u5c42\u662f\u5426\u5b58\u5728": adopted.get("GDB\u56fe\u5c42\u662f\u5426\u5b58\u5728", False),
                "\u91c7\u7528\u8bb0\u5f55GDB\u8def\u5f84": adopted.get("GDB\u8def\u5f84", ""),
                "\u5f53\u65e5\u8bb0\u5f55\u6b21\u6570": len(group),
                "\u5f53\u65e5\u6709\u6548\u6b21\u6570": int(valid_mask.sum()),
                "\u5f53\u65e5\u65e0\u6548\u6b21\u6570": int((~valid_mask).sum()),
            }
        )

    return pd.DataFrame(rows).sort_values(["\u6e56\u6cca\u7f16\u53f7", "\u65e5\u671f"]).reset_index(drop=True)


S2_REAL_IMAGE_DIR = PROJECT_ROOT / "data1" / "S2_22WEV_2024-06-01_to_2024-08-30"


def _list_real_s2_images() -> list[dict[str, object]]:
    items: list[dict[str, object]] = []
    if not S2_REAL_IMAGE_DIR.exists():
        return items
    pattern = re.compile(r"^S2_22WEV_(\d{8})T(\d{6})_T22WEV\.tif$", re.IGNORECASE)
    for tif_path in sorted(S2_REAL_IMAGE_DIR.glob("*.tif")):
        match = pattern.match(tif_path.name)
        if not match:
            continue
        date_str, time_str = match.groups()
        dt = pd.to_datetime(f"{date_str}{time_str}", format="%Y%m%d%H%M%S", errors="coerce")
        if pd.isna(dt):
            continue
        items.append(
            {
                "image_name": tif_path.name,
                "image_path": str(tif_path),
                "date": date_str,
                "time": time_str,
                "scene_key": f"{date_str}_{time_str}_S2X",
                "image_datetime": dt.strftime("%Y-%m-%d %H:%M"),
            }
        )
    return items


def _load_s2_preview_status_maps(selected_lake_ids: Iterable[str], lake_meta: pd.DataFrame | None = None) -> tuple[dict[tuple[str, str], str], dict[tuple[str, str], str]]:
    meta_lookup: dict[str, dict[str, object]] = {}
    if lake_meta is not None and not lake_meta.empty:
        meta_df = lake_meta.copy()
        meta_df["lake_id"] = meta_df["lake_id"].astype(str)
        meta_lookup = meta_df.set_index("lake_id").to_dict("index")

    status_map: dict[tuple[str, str], str] = {}
    preview_name_map: dict[tuple[str, str], str] = {}
    for lake_id in [str(v) for v in selected_lake_ids]:
        lake_meta_row = meta_lookup.get(lake_id, {})
        lake_type = str(lake_meta_row.get("lake_type") or lake_meta_row.get("size_category") or lake_meta_row.get("lake_size") or "")
        base_dir = S2_PREVIEW_ROOT / lake_type / lake_id / "01_Original_Base"
        clean_dir = S2_PREVIEW_ROOT / lake_type / lake_id / "02_Clean_Valid"
        cloud_dir = S2_PREVIEW_ROOT / lake_type / lake_id / "03_Cloud_Polluted"

        orig_keys = {}
        clean_keys = {}
        cloud_keys = {}
        if base_dir.exists():
            for p in base_dir.glob("*.png"):
                parts = _preview_scene_parts(p.name)
                if parts is not None:
                    orig_keys[parts[3]] = p.name
        if clean_dir.exists():
            for p in clean_dir.glob("*.png"):
                parts = _preview_scene_parts(p.name)
                if parts is not None:
                    clean_keys[parts[3]] = p.name
        if cloud_dir.exists():
            for p in cloud_dir.glob("*.png"):
                parts = _preview_scene_parts(p.name)
                if parts is not None:
                    cloud_keys[parts[3]] = p.name

        for scene_key, preview_name in orig_keys.items():
            key = (lake_id, scene_key)
            if scene_key in cloud_keys:
                status_map[key] = "cloud_covered"
                preview_name_map[key] = cloud_keys[scene_key]
            elif scene_key in clean_keys:
                status_map[key] = "normal"
                preview_name_map[key] = clean_keys[scene_key]
            else:
                status_map[key] = "normal"
                preview_name_map[key] = preview_name
    return status_map, preview_name_map


def build_s2_reference_status_records(selected_lake_ids: Iterable[str], lake_meta: pd.DataFrame | None = None) -> pd.DataFrame:
    selected_ids = [str(v) for v in selected_lake_ids]
    images = _list_real_s2_images()
    if not images:
        return pd.DataFrame()

    meta_lookup: dict[str, dict[str, object]] = {}
    if lake_meta is not None and not lake_meta.empty:
        meta_df = lake_meta.copy()
        meta_df["lake_id"] = meta_df["lake_id"].astype(str)
        meta_lookup = meta_df.set_index("lake_id").to_dict("index")

    status_map, _ = _load_s2_preview_status_maps(selected_ids, lake_meta)

    effective_df = _load_s2_effective_index()
    effective_lookup: dict[tuple[str, str], pd.Series] = {}
    effective_scene_set: set[tuple[str, str]] = set()
    if not effective_df.empty:
        effective_df = effective_df[effective_df["lake_id"].isin(selected_ids)].copy()
        for _, row in effective_df.iterrows():
            preview_name = str(row.get("preview_name", "") or "")
            parts = _preview_scene_parts(preview_name)
            if parts is None:
                continue
            scene_key = parts[3]
            lake_id = str(row["lake_id"])
            effective_lookup[(lake_id, scene_key)] = row
            effective_scene_set.add((lake_id, scene_key))

    rows: list[dict[str, object]] = []
    for lake_id in selected_ids:
        lake_meta_row = meta_lookup.get(lake_id, {})
        for item in images:
            scene_key = str(item["scene_key"])
            key = (lake_id, scene_key)
            effective_row = effective_lookup.get(key)
            rows.append(
                {
                    "\u6e56\u6cca\u7f16\u53f7": int(lake_id),
                    "\u6e56\u6cca\u9762\u79ef_\u5e73\u65b9\u7c73": lake_meta_row.get("area"),
                    "\u5f71\u50cf\u6587\u4ef6\u540d": item["image_name"],
                    "\u5f71\u50cf\u65e5\u671f": item["date"],
                    "\u5f71\u50cf\u65f6\u95f4": item["image_datetime"],
                    "\u573a\u666f\u952e": scene_key,
                    "\u6c34\u671f\u72b6\u6001": _water_status_from_fixed_window(str(item["date"])),
                    "\u8fb9\u754c\u72b6\u6001": _boundary_status_from_row(effective_row) or ("\u6709\u6548\u8fb9\u754c" if key in effective_scene_set else "\u975e\u6709\u6548\u8fb9\u754c"),
                    "\u5f71\u50cf\u72b6\u6001": status_map.get(key, "invalid"),
                    "\u5f71\u50cf\u8def\u5f84": item["image_path"],
                }
            )

    out = pd.DataFrame(rows)
    if out.empty:
        return out
    return out.sort_values(["\u6e56\u6cca\u7f16\u53f7", "\u5f71\u50cf\u65e5\u671f", "\u5f71\u50cf\u6587\u4ef6\u540d"]).reset_index(drop=True)
